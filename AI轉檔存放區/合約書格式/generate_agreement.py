#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Generate a formal Traditional Chinese property-sale proceeds agreement."""

from __future__ import annotations

import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Twips
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm, mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.pdfmetrics import registerFontFamily
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    FrameBreak,
    KeepTogether,
    NextPageTemplate,
    PageBreak,
    PageTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    HRFlowable,
    ListFlowable,
    ListItem,
)

OUT_DIR = os.path.dirname(os.path.abspath(__file__))
PDF_PATH = os.path.join(OUT_DIR, "房屋及土地出售價款分配同意書.pdf")
DOCX_PATH = os.path.join(OUT_DIR, "房屋及土地出售價款分配同意書.docx")
HTML_PATH = os.path.join(OUT_DIR, "房屋及土地出售價款分配同意書.html")

UMING = "/usr/share/fonts/truetype/arphic/uming.ttc"
UKAI = "/usr/share/fonts/truetype/arphic/ukai.ttc"
TW_INDEX = 2  # AR PL UMing TW / AR PL UKai TW

ADDRESS = "宜蘭縣頭城鎮拔雅里復興路12巷21號"
AGENCY = "永慶不動產羅東文化盛群加盟店"
AGENT = "郭火旺"
ELDER = "林天生"

BLANK = "＿"
def b(n: int = 10) -> str:
    return BLANK * n


def register_fonts() -> None:
    pdfmetrics.registerFont(TTFont("Ming", UMING, subfontIndex=TW_INDEX))
    pdfmetrics.registerFont(TTFont("Kai", UKAI, subfontIndex=TW_INDEX))
    registerFontFamily("Ming", normal="Ming", bold="Kai", italic="Ming", boldItalic="Kai")
    registerFontFamily("Kai", normal="Kai", bold="Kai", italic="Kai", boldItalic="Kai")


def make_styles() -> dict[str, ParagraphStyle]:
    justify = dict(alignment=TA_JUSTIFY, wordWrap="CJK", leading=22, firstLineIndent=0)
    return {
        "title": ParagraphStyle(
            "title",
            fontName="Kai",
            fontSize=22,
            leading=32,
            alignment=TA_CENTER,
            spaceAfter=4,
            textColor=colors.HexColor("#1a1a1a"),
        ),
        "subtitle": ParagraphStyle(
            "subtitle",
            fontName="Ming",
            fontSize=11,
            leading=18,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#333333"),
            spaceAfter=2,
        ),
        "meta": ParagraphStyle(
            "meta",
            fontName="Ming",
            fontSize=10,
            leading=16,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#444444"),
        ),
        "preamble_label": ParagraphStyle(
            "preamble_label",
            fontName="Kai",
            fontSize=12.5,
            leading=22,
            alignment=TA_LEFT,
            spaceBefore=8,
            spaceAfter=2,
        ),
        "body": ParagraphStyle(
            "body",
            fontName="Ming",
            fontSize=12,
            **justify,
            spaceBefore=2,
            spaceAfter=4,
        ),
        "body_indent": ParagraphStyle(
            "body_indent",
            fontName="Ming",
            fontSize=12,
            **justify,
            leftIndent=22,
            spaceBefore=1,
            spaceAfter=2,
        ),
        "article": ParagraphStyle(
            "article",
            fontName="Kai",
            fontSize=13,
            leading=24,
            alignment=TA_LEFT,
            spaceBefore=11,
            spaceAfter=4,
            textColor=colors.HexColor("#111111"),
        ),
        "item": ParagraphStyle(
            "item",
            fontName="Ming",
            fontSize=12,
            **justify,
            leftIndent=12,
            spaceBefore=2,
            spaceAfter=2,
        ),
        "subitem": ParagraphStyle(
            "subitem",
            fontName="Ming",
            fontSize=12,
            **justify,
            leftIndent=28,
            spaceBefore=1,
            spaceAfter=1,
        ),
        "note": ParagraphStyle(
            "note",
            fontName="Ming",
            fontSize=10,
            leading=16,
            alignment=TA_JUSTIFY,
            wordWrap="CJK",
            textColor=colors.HexColor("#333333"),
            leftIndent=12,
            spaceBefore=2,
            spaceAfter=4,
        ),
        "annex_title": ParagraphStyle(
            "annex_title",
            fontName="Kai",
            fontSize=16,
            leading=26,
            alignment=TA_CENTER,
            spaceBefore=6,
            spaceAfter=10,
        ),
        "sign_title": ParagraphStyle(
            "sign_title",
            fontName="Kai",
            fontSize=13,
            leading=22,
            alignment=TA_LEFT,
            spaceBefore=10,
            spaceAfter=6,
        ),
        "sign_cell": ParagraphStyle(
            "sign_cell",
            fontName="Ming",
            fontSize=10,
            leading=15,
            alignment=TA_LEFT,
            wordWrap="CJK",
        ),
        "sign_cell_c": ParagraphStyle(
            "sign_cell_c",
            fontName="Ming",
            fontSize=10,
            leading=15,
            alignment=TA_CENTER,
            wordWrap="CJK",
        ),
        "table_head": ParagraphStyle(
            "table_head",
            fontName="Kai",
            fontSize=10,
            leading=14,
            alignment=TA_CENTER,
        ),
        "table_cell": ParagraphStyle(
            "table_cell",
            fontName="Ming",
            fontSize=10,
            leading=14,
            alignment=TA_CENTER,
            wordWrap="CJK",
        ),
        "table_cell_l": ParagraphStyle(
            "table_cell_l",
            fontName="Ming",
            fontSize=10,
            leading=14,
            alignment=TA_LEFT,
            wordWrap="CJK",
        ),
        "date": ParagraphStyle(
            "date",
            fontName="Kai",
            fontSize=13,
            leading=22,
            alignment=TA_CENTER,
            spaceBefore=16,
            spaceAfter=4,
        ),
        "footer": ParagraphStyle(
            "footer",
            fontName="Ming",
            fontSize=8.5,
            leading=12,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#555555"),
        ),
        "small_center": ParagraphStyle(
            "small_center",
            fontName="Ming",
            fontSize=9,
            leading=14,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#444444"),
        ),
    }


def header_footer(canvas, doc) -> None:
    canvas.saveState()
    w, h = A4
    canvas.setStrokeColor(colors.HexColor("#222222"))
    canvas.setLineWidth(1.1)
    canvas.line(1.8 * cm, h - 1.45 * cm, w - 1.8 * cm, h - 1.45 * cm)
    canvas.setLineWidth(0.35)
    canvas.line(1.8 * cm, h - 1.58 * cm, w - 1.8 * cm, h - 1.58 * cm)
    canvas.setFont("Kai", 9)
    canvas.setFillColor(colors.HexColor("#333333"))
    canvas.drawString(1.8 * cm, h - 1.32 * cm, "房屋及土地出售價款分配同意書")
    canvas.setFont("Ming", 8.5)
    canvas.drawRightString(w - 1.8 * cm, h - 1.32 * cm, ADDRESS)

    canvas.setLineWidth(0.35)
    canvas.line(1.8 * cm, 1.45 * cm, w - 1.8 * cm, 1.45 * cm)
    canvas.setLineWidth(1.1)
    canvas.line(1.8 * cm, 1.32 * cm, w - 1.8 * cm, 1.32 * cm)
    canvas.setFont("Ming", 8.5)
    canvas.setFillColor(colors.HexColor("#444444"))
    page = canvas.getPageNumber()
    canvas.drawCentredString(w / 2, 0.92 * cm, f"第 {page} 頁")
    canvas.drawString(1.8 * cm, 0.92 * cm, "簽署後請分送共有人、代書及履約保證機構各執乙份")
    canvas.restoreState()


def first_page(canvas, doc) -> None:
    canvas.saveState()
    w, h = A4
    canvas.setStrokeColor(colors.HexColor("#222222"))
    canvas.setLineWidth(0.35)
    canvas.line(1.8 * cm, 1.45 * cm, w - 1.8 * cm, 1.45 * cm)
    canvas.setLineWidth(1.1)
    canvas.line(1.8 * cm, 1.32 * cm, w - 1.8 * cm, 1.32 * cm)
    canvas.setFont("Ming", 8.5)
    canvas.setFillColor(colors.HexColor("#444444"))
    page = canvas.getPageNumber()
    canvas.drawCentredString(w / 2, 0.92 * cm, f"第 {page} 頁")
    canvas.restoreState()


def P(text: str, style: ParagraphStyle) -> Paragraph:
    return Paragraph(text.replace("\n", "<br/>"), style)


def checkbox(label: str, styles) -> Paragraph:
    return P(f"□　{label}", styles["item"])


def article_block(title: str, parts: list, styles) -> KeepTogether:
    flow = [P(title, styles["article"])]
    flow.extend(parts)
    return KeepTogether(flow)


def signature_block(name: str, styles, pref_note: str = "") -> Table:
    note = f"<font color='#555555'>（{pref_note}）</font>" if pref_note else ""
    name_line = f"姓名：{name or b(12)}{note}"
    data = [
        [
            P(name_line, styles["sign_cell"]),
            P(f"國民身分證統一編號：{b(14)}", styles["sign_cell"]),
        ],
        [
            P(f"戶籍地址：{b(28)}", styles["sign_cell"]),
            P(f"聯絡電話：{b(14)}", styles["sign_cell"]),
        ],
        [
            P(f"權利範圍（持分）：{b(12)}", styles["sign_cell"]),
            P("簽名或蓋章：", styles["sign_cell"]),
        ],
        [
            P(f"指定撥款帳戶（銀行／帳號／戶名）：{b(18)}", styles["sign_cell"]),
            P("（請蓋印鑑章，與印鑑證明相符）", styles["sign_cell"]),
        ],
    ]
    t = Table(data, colWidths=[9.2 * cm, 7.4 * cm], rowHeights=[1.05 * cm, 1.05 * cm, 1.15 * cm, 1.15 * cm])
    t.setStyle(
        TableStyle(
            [
                ("BOX", (0, 0), (-1, -1), 0.6, colors.HexColor("#222222")),
                ("INNERGRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#888888")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f4f1ea")),
                ("SPAN", (0, 3), (0, 3)),
            ]
        )
    )
    return t


def annex_table(headers, rows, col_widths, styles, row_h=1.15 * cm) -> Table:
    head = [P(h, styles["table_head"]) for h in headers]
    body = []
    for row in rows:
        body.append([P(c, styles["table_cell"] if i else styles["table_cell"]) for i, c in enumerate(row)])
    t = Table([head] + body, colWidths=col_widths, rowHeights=[0.85 * cm] + [row_h] * len(rows))
    t.setStyle(
        TableStyle(
            [
                ("BOX", (0, 0), (-1, -1), 0.8, colors.HexColor("#222222")),
                ("INNERGRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#555555")),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2c2c2c")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    # header needs white text - Paragraph uses its own color; rebuild header in white
    return t


def annex_table_white_header(headers, rows, col_widths, styles, row_h=1.2 * cm) -> Table:
    head_style = ParagraphStyle(
        "th_white",
        parent=styles["table_head"],
        textColor=colors.white,
        fontName="Kai",
        fontSize=10,
    )
    head = [P(h, head_style) for h in headers]
    body = [[P(c, styles["table_cell"]) for c in row] for row in rows]
    t = Table([head] + body, colWidths=col_widths, rowHeights=[0.9 * cm] + [row_h] * len(rows), repeatRows=1)
    t.setStyle(
        TableStyle(
            [
                ("BOX", (0, 0), (-1, -1), 0.8, colors.HexColor("#222222")),
                ("INNERGRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#666666")),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2a2a2a")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    return t


def build_story(styles) -> list:
    s = styles
    story = []

    story.append(Spacer(1, 6))
    story.append(P("房屋及土地出售價款分配同意書", s["title"]))
    story.append(HRFlowable(width="100%", thickness=1.6, color=colors.HexColor("#222222"), spaceBefore=4, spaceAfter=1))
    story.append(HRFlowable(width="100%", thickness=0.4, color=colors.HexColor("#222222"), spaceBefore=0, spaceAfter=8))
    story.append(P("（全體房屋所有權共有人簽署）", s["subtitle"]))
    story.append(P(f"買賣標的門牌：{ADDRESS}", s["meta"]))
    story.append(Spacer(1, 8))

    story.append(P("立同意書人", s["preamble_label"]))
    story.append(
        P(
            "本件買賣標的（定義如第一條）之<b>全體房屋及土地所有權共有人</b>"
            "（以下合稱「全體共有人」或「立同意書人」；各人之姓名、國民身分證統一編號、"
            "戶籍地址、權利範圍及指定撥款帳戶，詳如本同意書末之簽署欄及附件）。",
            s["body"],
        )
    )
    story.append(
        P(
            "緣立同意書人等為第一條所載房地之登記所有權人，為使出售程序明確、價款分配公允、"
            "交割撥款有據，經全體協商後，<b>本於自由意願</b>，訂立本同意書，以資共同遵守：",
            s["body"],
        )
    )

    story.append(
        article_block(
            "第一條　買賣標的",
            [
                P("本同意書所稱「買賣標的」，係指下列房屋及其基地，並包含依法應隨同移轉之權利：", s["body"]),
                P(f"（一）門牌地址：<b>{ADDRESS}</b>。", s["item"]),
                P(
                    f"（二）土地標示：{b(4)}縣（市）{b(6)}鄉（鎮、市、區）{b(6)}段{b(4)}小段{b(6)}地號；"
                    f"地目{b(4)}；面積{b(8)}平方公尺；權利範圍{b(10)}。",
                    s["item"],
                ),
                P(
                    f"（三）建物標示：建號{b(10)}；門牌同第（一）款；建築完成日期{b(10)}；"
                    f"面積{b(8)}平方公尺（含附屬建物）；權利範圍{b(10)}。",
                    s["item"],
                ),
                P(
                    "前項土地及建物之詳細標示，如與地政機關最新登記謄本不符，以登記謄本為準；"
                    "得於空白處補正或檢附謄本作為附件，視為本同意書之一部分。",
                    s["body"],
                ),
            ],
            s,
        )
    )

    story.append(
        article_block(
            "第二條　委託銷售",
            [
                P(
                    f"全體共有人同意將買賣標的委託<b>{AGENCY}</b>辦理銷售，"
                    f"並指定業務代表<b>{AGENT}先生</b>執行帶看、議價、媒合及簽約協商等銷售業務。",
                    s["body"],
                ),
                P(
                    "委託期間、仲介服務報酬及其他權利義務，另以全體共有人與該加盟店簽訂之"
                    "不動產委託銷售契約為準。該契約與本同意書抵觸時，<b>有關價款分配及撥款指示，以本同意書為準</b>。",
                    s["body"],
                ),
            ],
            s,
        )
    )

    story.append(
        article_block(
            "第三條　委託代書辦理交割",
            [
                P(
                    "全體共有人同意本件買賣之簽約、用印、稅費申報、所有權移轉登記、點交、"
                    "抵押權塗銷（如有）及其他交割程序，委由代書依相關法令、買賣契約及本同意書辦理。",
                    s["body"],
                ),
                P(f"受託代書／事務所名稱：{b(18)}　　聯絡電話：{b(12)}", s["item"]),
                P(f"履約保證機構名稱：{b(18)}　　專案編號：{b(12)}", s["item"]),
            ],
            s,
        )
    )

    story.append(
        article_block(
            "第四條　出售價金及應扣除費用",
            [
                P(
                    "一、本次出售總價金，以全體共有人與買方簽訂之不動產買賣契約所載價款為準（下稱「成交總價」）。"
                    f"預定或已議定之成交總價：新臺幣{b(14)}元整（未填者，以日後簽訂之買賣契約為準）。",
                    s["item"],
                ),
                P(
                    "二、成交總價應優先扣除下列費用及稅捐後，以其餘額作為本同意書之可分配金額（下稱「可分配價金」）：",
                    s["item"],
                ),
                P("（一）仲介服務報酬；", s["subitem"]),
                P(
                    "（二）土地增值稅、依法應由出賣人負擔之房屋稅及地價稅、印花稅、"
                    "房地合一所得稅，以及其他依法應由出賣人負擔之稅捐；",
                    s["subitem"],
                ),
                P(
                    "（三）代書費、地政規費、謄本費、印鑑證明及印鑑登記相關費用、抵押權塗銷費、履約保證相關費用；",
                    s["subitem"],
                ),
                P("（四）為完成本件交易必要，且經全體共有人同意或依法必須支出之其他費用。", s["subitem"]),
                P(
                    "三、前項費用之實際金額，以代書製作之結算明細、繳款收據或稅單為準。"
                    "全體共有人簽署本同意書，即授權代書據實扣除後，依第五條及附件分配撥款。",
                    s["item"],
                ),
            ],
            s,
        )
    )

    story.append(
        article_block(
            "第五條　價款分配方式",
            [
                P(
                    f"一、立同意書人<b>{ELDER}</b>（於兄弟姊妹間為最年長之兄長），基於親屬間協商之結果，"
                    "同意自可分配價金中領取下列方式之一（請擇一勾選並以正楷填載；"
                    "勾選後之空白處如未填載，本項不生分配效力，應由全體共有人補簽後始得撥款）：",
                    s["item"],
                ),
                checkbox(
                    f"可分配價金之百分之{b(6)}（即　{b(4)}　％；大寫：百分之{b(8)}）。",
                    s,
                ),
                checkbox(
                    f"固定金額新臺幣{b(14)}元整（大寫：新臺幣{b(16)}元整）。",
                    s,
                ),
                P(
                    f"二、可分配價金扣除{ELDER}依前項取得之數額後，<b>其餘價款由其他所有權共有人協議分配</b>，"
                    "其姓名、比例或金額，詳如附件一「其他共有人價款分配明細表」。",
                    s["item"],
                ),
                P(
                    "三、全體共有人確認並同意：本條約定之分配方式，係本於兄弟姊妹及共有人間之特別合意，"
                    "<b>得與地政登記之應有部分（持分）比例不相同</b>；一經簽署，"
                    "不得再以登記持分、出資多寡、管理貢獻或其他理由異議，或請求重新分配。",
                    s["item"],
                ),
                P(
                    "四、如成交總價或應扣除費用於簽約後有增減，致可分配價金變動時："
                    f"勾選「百分比」者，應按原約定比例自動調整；勾選「固定金額」者，{ELDER}所領固定金額不變，"
                    "其餘共有人就剩餘款項再協議，並以書面補正附件一。",
                    s["item"],
                ),
            ],
            s,
        )
    )

    story.append(
        article_block(
            "第六條　撥款指示",
            [
                P(
                    "一、全體共有人在此同意並指示受託代書及履約保證機構，於買賣價金依約可撥付時，"
                    "<b>應逕依本同意書第五條、附件一及附件二辦理分配撥款</b>，無需再逐一徵詢各共有人。"
                    "各方不得對依本同意書所為之撥款提出異議。",
                    s["item"],
                ),
                P(
                    "二、各共有人指定之撥款帳戶詳如簽署欄及附件二。帳戶資料如有變更，"
                    "應於撥款日至少三個營業日前，以書面通知代書及履約保證機構，並由變更人簽名或蓋章。",
                    s["item"],
                ),
                P(
                    "三、因帳戶資料錯誤、未及時變更、拒絕提供證件印鑑，或受款帳戶非其本人名義，"
                    "致撥款遲延、退匯或誤入他人帳戶者，由該共有人自行負責，不得因此拒絕交割或主張本同意書無效。",
                    s["item"],
                ),
                P(
                    "四、代書及履約保證機構依本同意書完成撥款後，就該部分價金對全體共有人之給付義務即為清償完畢。",
                    s["item"],
                ),
            ],
            s,
        )
    )

    story.append(
        article_block(
            "第七條　聲明、保證及效力",
            [
                P(
                    "一、立同意書人簽署本同意書，均出於自由意願，並已詳細閱讀、充分瞭解全部內容，並願受其拘束。",
                    s["item"],
                ),
                P(
                    "二、立同意書人保證其為買賣標的之合法共有人（或有權處分之人），"
                    "並已將所知之其他共有人、他項權利、優先購買權、租賃、占用或其他可能影響交易之事由，"
                    "據實告知受託代書及仲介。",
                    s["item"],
                ),
                P(
                    "三、如因立同意書人隱匿、虛偽陳述或未配合提供權狀、印鑑證明、身分證件，"
                    "致本件交易或他方受有損害，應由可歸責之立同意書人負損害賠償責任。",
                    s["item"],
                ),
                P(
                    "四、本同意書如有塗改，應於塗改處由全體共有人蓋章，始生效力。空白未填部分，不得擅自添註。",
                    s["item"],
                ),
            ],
            s,
        )
    )

    story.append(
        article_block(
            "第八條　補充約定、份數、準據法及生效",
            [
                P(
                    "一、本同意書未盡事宜，依中華民國民法、土地法、土地登記規則及其他相關法令辦理；"
                    "必要時得由全體共有人另行簽立書面協議補充，該書面與本同意書有同等效力。",
                    s["item"],
                ),
                P(
                    "二、本同意書之解釋、效力與履行，以中華民國法律為準據法。"
                    "如因此發生訴訟，同意以臺灣宜蘭地方法院為第一審管轄法院。",
                    s["item"],
                ),
                P(
                    f"三、本同意書壹式　{b(4)}　份，由全體共有人、受託代書、履約保證機構及受託仲介店各執乙份為憑；"
                    "影本或掃描檔經核與正本無異者，與正本有同一效力。",
                    s["item"],
                ),
                P("四、本同意書於全體共有人親自簽名或蓋用與印鑑證明相符之印鑑章之日起生效。", s["item"]),
            ],
            s,
        )
    )

    story.append(Spacer(1, 10))
    story.append(P("此　　據", s["date"]))
    story.append(P("中華民國　　　　年　　　　月　　　　日", s["date"]))

    story.append(PageBreak())
    story.append(P("立同意書人簽署欄", s["annex_title"]))
    story.append(
        P(
            "以下簽署人確為本同意書之立同意書人，已閱讀並同意全文及附件，願依約履行。"
            "未使用之簽署欄，請劃記「空白」並由在場共有人蓋章註銷。",
            s["small_center"],
        )
    )
    story.append(Spacer(1, 8))

    story.append(signature_block(ELDER, s, "長兄／共有人"))
    story.append(Spacer(1, 7))
    for i in range(2, 9):
        story.append(signature_block("", s, f"共有人 {i}"))
        story.append(Spacer(1, 7))

    story.append(Spacer(1, 4))
    story.append(P("見證及收執確認", s["sign_title"]))

    witness = [
        [
            P("<b>受託仲介</b>", s["table_head"]),
            P("<b>受託代書</b>", s["table_head"]),
            P("<b>履約保證機構</b>", s["table_head"]),
        ],
        [
            P(
                f"{AGENCY}<br/>業務代表：{AGENT}<br/><br/>簽署／店章：<br/><br/><br/>日期：　　年　　月　　日",
                s["sign_cell"],
            ),
            P(
                f"事務所：{b(12)}<br/>代書：{b(10)}<br/><br/>簽署／職章：<br/><br/><br/>日期：　　年　　月　　日",
                s["sign_cell"],
            ),
            P(
                f"機構：{b(12)}<br/>承辦：{b(10)}<br/><br/>簽署／職章：<br/><br/><br/>日期：　　年　　月　　日",
                s["sign_cell"],
            ),
        ],
    ]
    wt = Table(witness, colWidths=[5.55 * cm, 5.55 * cm, 5.55 * cm], rowHeights=[0.8 * cm, 4.2 * cm])
    wt.setStyle(
        TableStyle(
            [
                ("BOX", (0, 0), (-1, -1), 0.8, colors.HexColor("#222222")),
                ("INNERGRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#666666")),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#efefe9")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    story.append(wt)

    story.append(PageBreak())
    story.append(P("附件一　其他共有人價款分配明細表", s["annex_title"]))
    story.append(
        P(
            f"本表係第五條第二項之分配依據。可分配價金扣除立同意書人{ELDER}依第五條第一項取得之數額後，"
            "其餘額按下表分配。比例合計應為百分之一百。金額欄得於結算後由代書補填。",
            s["body"],
        )
    )
    story.append(Spacer(1, 8))

    annex1_rows = []
    annex1_rows.append([ELDER, "依第五條第一項（百分比或固定金額）", "依第五條第一項計算", ""])
    for i in range(1, 8):
        annex1_rows.append([b(8), f"{b(6)}％　或　新臺幣{b(8)}元", b(10), ""])
    annex1_rows.append(["合計", "100％（其餘額部分）", "可分配價金全額", ""])

    story.append(
        annex_table_white_header(
            ["共有人姓名", "分配比例或金額", "預估／實際分配金額", "簽認"],
            annex1_rows,
            [3.6 * cm, 5.6 * cm, 4.4 * cm, 3.0 * cm],
            s,
            row_h=1.05 * cm,
        )
    )
    story.append(Spacer(1, 8))
    story.append(
        P(
            "備註：如實際可分配價金與簽約時預估不同，除第五條第四項另有約定外，按本表比例調整。"
            "本表塗改處應由全體共有人蓋章。",
            s["note"],
        )
    )
    story.append(P(f"代書結算日期：中華民國{b(4)}年{b(3)}月{b(3)}日　　代書簽章：{b(10)}", s["item"]))

    story.append(PageBreak())
    story.append(P("附件二　指定撥款帳戶及應備文件", s["annex_title"]))
    story.append(
        P(
            "請以共有人本人名義帳戶為原則。非本人帳戶應另附同意書。撥款前請核對戶名、銀行及帳號無誤。",
            s["body"],
        )
    )
    story.append(Spacer(1, 8))

    annex2_rows = []
    annex2_rows.append([ELDER, b(10), b(12), b(10), ""])
    for _ in range(7):
        annex2_rows.append([b(8), b(10), b(12), b(10), ""])

    story.append(
        annex_table_white_header(
            ["戶名（共有人）", "金融機構／分行", "帳號", "身分證字號", "簽認"],
            annex2_rows,
            [3.3 * cm, 4.2 * cm, 3.8 * cm, 3.0 * cm, 2.3 * cm],
            s,
            row_h=1.05 * cm,
        )
    )
    story.append(Spacer(1, 12))
    story.append(P("交割時請備文件（請代書勾選）", s["sign_title"]))
    checks = [
        "□　國民身分證正本及影本　　□　印鑑章及印鑑證明　　□　所有權狀正本",
        "□　委託書／同意書正本　　□　戶口名簿或戶籍謄本　　□　指定帳戶存摺封面影本",
        "□　他項權利塗銷相關文件（如有）　　□　其他：＿＿＿＿＿＿＿＿＿＿＿＿",
    ]
    for line in checks:
        story.append(P(line, s["item"]))

    story.append(Spacer(1, 18))
    story.append(HRFlowable(width="100%", thickness=0.4, color=colors.HexColor("#222222"), spaceBefore=4, spaceAfter=8))
    story.append(
        P(
            "本同意書係依立同意書人提供之協議意旨撰擬，供全體共有人簽署使用。"
            "涉及稅負、持分、優先購買權及繼承關係等事項，建議於用印前請受託代書或律師核閱。",
            s["small_center"],
        )
    )
    story.append(P(f"標的門牌：{ADDRESS}　　受託仲介：{AGENCY}　業務代表：{AGENT}", s["small_center"]))

    return story


def build_pdf() -> None:
    register_fonts()
    styles = make_styles()
    doc = BaseDocTemplate(
        PDF_PATH,
        pagesize=A4,
        leftMargin=1.8 * cm,
        rightMargin=1.8 * cm,
        topMargin=1.9 * cm,
        bottomMargin=1.85 * cm,
        title="房屋及土地出售價款分配同意書",
        author="全體房屋所有權共有人",
        subject=ADDRESS,
    )
    frame_first = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height + 0.35 * cm, id="first")
    frame_later = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height - 0.25 * cm, id="later")
    doc.addPageTemplates(
        [
            PageTemplate(id="First", frames=frame_first, onPage=first_page),
            PageTemplate(id="Later", frames=frame_later, onPage=header_footer),
        ]
    )
    story = [NextPageTemplate("Later")]
    story.extend(build_story(styles))
    doc.build(story)


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------

def set_run_font(run, font: str, size_pt: float, bold: bool = False) -> None:
    run.bold = bold
    run.font.size = Pt(size_pt)
    run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
    run.font.name = font
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), font)
    rFonts.set(qn("w:hAnsi"), font)
    rFonts.set(qn("w:eastAsia"), font)
    rFonts.set(qn("w:cs"), font)


def set_paragraph_format(p, align="left", space_before=0, space_after=6, line=22, first_indent=None) -> None:
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = Pt(line)
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "justify":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if first_indent is not None:
        pf.first_line_indent = Cm(first_indent)


def add_text(p, text: str, font="新細明體", size=12, bold=False) -> None:
    run = p.add_run(text)
    set_run_font(run, font, size, bold)


def add_p(doc, text: str, font="新細明體", size=12, bold=False, align="justify", sb=2, sa=4, line=22, indent=None):
    p = doc.add_paragraph()
    set_paragraph_format(p, align=align, space_before=sb, space_after=sa, line=line)
    if indent:
        p.paragraph_format.left_indent = Cm(indent)
    add_text(p, text, font, size, bold)
    return p


def set_cell_border(cell, **kwargs) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), kwargs.get("val", "single"))
        element.set(qn("w:sz"), kwargs.get("sz", "8"))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), kwargs.get("color", "222222"))
        tcBorders.append(element)
    tcPr.append(tcBorders)


def shade_cell(cell, fill: str) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_text(cell, text: str, font="新細明體", size=10, bold=False, align="left") -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    add_text(p, text, font, size, bold)


def prevent_break(paragraph) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext")
    pPr.append(keep)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr)
    run._r.append(fld_char_end)
    set_run_font(run, "新細明體", 9)


def build_docx() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.9)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.different_first_page_header_footer = True

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(fp, "房屋及土地出售價款分配同意書　·　第 ", "新細明體", 9)
    add_page_number(fp)
    add_text(fp, " 頁　·　" + ADDRESS, "新細明體", 9)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(hp, "房屋及土地出售價款分配同意書", "標楷體", 10)

    # default east-asia font
    style = doc.styles["Normal"]
    style.font.name = "新細明體"
    style.font.size = Pt(12)
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), "新細明體")
    rFonts.set(qn("w:ascii"), "Times New Roman")

    add_p(doc, "房屋及土地出售價款分配同意書", font="標楷體", size=22, bold=True, align="center", sb=6, sa=4, line=32)
    add_p(doc, "（全體房屋所有權共有人簽署）", font="新細明體", size=11, align="center", sb=0, sa=2, line=18)
    add_p(doc, f"買賣標的門牌：{ADDRESS}", font="新細明體", size=11, align="center", sb=0, sa=10, line=18)

    add_p(doc, "立同意書人", font="標楷體", size=13, bold=True, align="left", sb=8, sa=4, line=22)
    add_p(
        doc,
        "本件買賣標的（定義如第一條）之全體房屋及土地所有權共有人（以下合稱「全體共有人」或「立同意書人」；"
        "各人之姓名、國民身分證統一編號、戶籍地址、權利範圍及指定撥款帳戶，詳如本同意書末之簽署欄及附件）。",
    )
    add_p(
        doc,
        "緣立同意書人等為第一條所載房地之登記所有權人，為使出售程序明確、價款分配公允、交割撥款有據，"
        "經全體協商後，本於自由意願，訂立本同意書，以資共同遵守：",
    )

    articles = [
        (
            "第一條　買賣標的",
            [
                "本同意書所稱「買賣標的」，係指下列房屋及其基地，並包含依法應隨同移轉之權利：",
                f"（一）門牌地址：{ADDRESS}。",
                "（二）土地標示：＿＿＿＿縣（市）＿＿＿＿＿＿鄉（鎮、市、區）＿＿＿＿＿＿段＿＿＿＿小段＿＿＿＿＿＿地號；"
                "地目＿＿＿＿；面積＿＿＿＿＿＿＿＿平方公尺；權利範圍＿＿＿＿＿＿＿＿＿＿。",
                "（三）建物標示：建號＿＿＿＿＿＿＿＿＿＿；門牌同第（一）款；建築完成日期＿＿＿＿＿＿＿＿＿＿；"
                "面積＿＿＿＿＿＿＿＿平方公尺（含附屬建物）；權利範圍＿＿＿＿＿＿＿＿＿＿。",
                "前項土地及建物之詳細標示，如與地政機關最新登記謄本不符，以登記謄本為準；"
                "得於空白處補正或檢附謄本作為附件，視為本同意書之一部分。",
            ],
        ),
        (
            "第二條　委託銷售",
            [
                f"全體共有人同意將買賣標的委託「{AGENCY}」辦理銷售，並指定業務代表{AGENT}先生執行帶看、議價、媒合及簽約協商等銷售業務。",
                "委託期間、仲介服務報酬及其他權利義務，另以全體共有人與該加盟店簽訂之不動產委託銷售契約為準。"
                "該契約與本同意書抵觸時，有關價款分配及撥款指示，以本同意書為準。",
            ],
        ),
        (
            "第三條　委託代書辦理交割",
            [
                "全體共有人同意本件買賣之簽約、用印、稅費申報、所有權移轉登記、點交、抵押權塗銷（如有）及其他交割程序，"
                "委由代書依相關法令、買賣契約及本同意書辦理。",
                "受託代書／事務所名稱：＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿　　聯絡電話：＿＿＿＿＿＿＿＿＿＿＿＿",
                "履約保證機構名稱：＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿　　專案編號：＿＿＿＿＿＿＿＿＿＿＿＿",
            ],
        ),
        (
            "第四條　出售價金及應扣除費用",
            [
                "一、本次出售總價金，以全體共有人與買方簽訂之不動產買賣契約所載價款為準（下稱「成交總價」）。"
                "預定或已議定之成交總價：新臺幣＿＿＿＿＿＿＿＿＿＿＿＿＿＿元整（未填者，以日後簽訂之買賣契約為準）。",
                "二、成交總價應優先扣除下列費用及稅捐後，以其餘額作為本同意書之可分配金額（下稱「可分配價金」）：",
                "（一）仲介服務報酬；",
                "（二）土地增值稅、依法應由出賣人負擔之房屋稅及地價稅、印花稅、房地合一所得稅，以及其他依法應由出賣人負擔之稅捐；",
                "（三）代書費、地政規費、謄本費、印鑑證明及印鑑登記相關費用、抵押權塗銷費、履約保證相關費用；",
                "（四）為完成本件交易必要，且經全體共有人同意或依法必須支出之其他費用。",
                "三、前項費用之實際金額，以代書製作之結算明細、繳款收據或稅單為準。"
                "全體共有人簽署本同意書，即授權代書據實扣除後，依第五條及附件分配撥款。",
            ],
        ),
        (
            "第五條　價款分配方式",
            [
                f"一、立同意書人{ELDER}（於兄弟姊妹間為最年長之兄長），基於親屬間協商之結果，"
                "同意自可分配價金中領取下列方式之一（請擇一勾選並以正楷填載；"
                "勾選後之空白處如未填載，本項不生分配效力，應由全體共有人補簽後始得撥款）：",
                "□　可分配價金之百分之＿＿＿＿＿＿（即　＿＿＿＿　％；大寫：百分之＿＿＿＿＿＿＿＿）。",
                "□　固定金額新臺幣＿＿＿＿＿＿＿＿＿＿＿＿＿＿元整（大寫：新臺幣＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿元整）。",
                f"二、可分配價金扣除{ELDER}依前項取得之數額後，其餘價款由其他所有權共有人協議分配，"
                "其姓名、比例或金額，詳如附件一「其他共有人價款分配明細表」。",
                "三、全體共有人確認並同意：本條約定之分配方式，係本於兄弟姊妹及共有人間之特別合意，"
                "得與地政登記之應有部分（持分）比例不相同；一經簽署，不得再以登記持分、出資多寡、管理貢獻或其他理由異議，或請求重新分配。",
                "四、如成交總價或應扣除費用於簽約後有增減，致可分配價金變動時："
                f"勾選「百分比」者，應按原約定比例自動調整；勾選「固定金額」者，{ELDER}所領固定金額不變，"
                "其餘共有人就剩餘款項再協議，並以書面補正附件一。",
            ],
        ),
        (
            "第六條　撥款指示",
            [
                "一、全體共有人在此同意並指示受託代書及履約保證機構，於買賣價金依約可撥付時，"
                "應逕依本同意書第五條、附件一及附件二辦理分配撥款，無需再逐一徵詢各共有人。"
                "各方不得對依本同意書所為之撥款提出異議。",
                "二、各共有人指定之撥款帳戶詳如簽署欄及附件二。帳戶資料如有變更，"
                "應於撥款日至少三個營業日前，以書面通知代書及履約保證機構，並由變更人簽名或蓋章。",
                "三、因帳戶資料錯誤、未及時變更、拒絕提供證件印鑑，或受款帳戶非其本人名義，"
                "致撥款遲延、退匯或誤入他人帳戶者，由該共有人自行負責，不得因此拒絕交割或主張本同意書無效。",
                "四、代書及履約保證機構依本同意書完成撥款後，就該部分價金對全體共有人之給付義務即為清償完畢。",
            ],
        ),
        (
            "第七條　聲明、保證及效力",
            [
                "一、立同意書人簽署本同意書，均出於自由意願，並已詳細閱讀、充分瞭解全部內容，並願受其拘束。",
                "二、立同意書人保證其為買賣標的之合法共有人（或有權處分之人），"
                "並已將所知之其他共有人、他項權利、優先購買權、租賃、占用或其他可能影響交易之事由，據實告知受託代書及仲介。",
                "三、如因立同意書人隱匿、虛偽陳述或未配合提供權狀、印鑑證明、身分證件，"
                "致本件交易或他方受有損害，應由可歸責之立同意書人負損害賠償責任。",
                "四、本同意書如有塗改，應於塗改處由全體共有人蓋章，始生效力。空白未填部分，不得擅自添註。",
            ],
        ),
        (
            "第八條　補充約定、份數、準據法及生效",
            [
                "一、本同意書未盡事宜，依中華民國民法、土地法、土地登記規則及其他相關法令辦理；"
                "必要時得由全體共有人另行簽立書面協議補充，該書面與本同意書有同等效力。",
                "二、本同意書之解釋、效力與履行，以中華民國法律為準據法。如因此發生訴訟，同意以臺灣宜蘭地方法院為第一審管轄法院。",
                "三、本同意書壹式　＿＿＿＿　份，由全體共有人、受託代書、履約保證機構及受託仲介店各執乙份為憑；"
                "影本或掃描檔經核與正本無異者，與正本有同一效力。",
                "四、本同意書於全體共有人親自簽名或蓋用與印鑑證明相符之印鑑章之日起生效。",
            ],
        ),
    ]

    for title, paras in articles:
        add_p(doc, title, font="標楷體", size=13, bold=True, align="left", sb=12, sa=6, line=24)
        for t in paras:
            add_p(doc, t, indent=0.2 if t.startswith(("一", "二", "三", "四", "□", "（")) else 0)

    add_p(doc, "此　　據", font="標楷體", size=13, align="center", sb=16, sa=8, line=24)
    add_p(doc, "中華民國　　　　年　　　　月　　　　日", font="標楷體", size=13, align="center", sb=4, sa=12, line=24)

    doc.add_page_break()
    add_p(doc, "立同意書人簽署欄", font="標楷體", size=16, bold=True, align="center", sb=4, sa=6, line=26)
    add_p(
        doc,
        "以下簽署人確為本同意書之立同意書人，已閱讀並同意全文及附件，願依約履行。"
        "未使用之簽署欄，請劃記「空白」並由在場共有人蓋章註銷。",
        size=10,
        align="center",
        sa=10,
        line=16,
    )

    names = [(ELDER, "長兄／共有人")] + [("", f"共有人 {i}") for i in range(2, 9)]
    for name, note in names:
        display = f"{name}（{note}）" if name else f"{'＿'*12}（{note}）"
        table = doc.add_table(rows=4, cols=2)
        table.autofit = True
        table.allow_autofit = True
        cells = [
            (0, 0, f"姓名：{display}"),
            (0, 1, "國民身分證統一編號：＿＿＿＿＿＿＿＿＿＿＿＿＿＿"),
            (1, 0, "戶籍地址：＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿"),
            (1, 1, "聯絡電話：＿＿＿＿＿＿＿＿＿＿＿＿＿＿"),
            (2, 0, "權利範圍（持分）：＿＿＿＿＿＿＿＿＿＿＿＿"),
            (2, 1, "簽名或蓋章："),
            (3, 0, "指定撥款帳戶（銀行／帳號／戶名）：＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿"),
            (3, 1, "（請蓋印鑑章，與印鑑證明相符）"),
        ]
        for r, c, text in cells:
            cell = table.cell(r, c)
            set_cell_text(cell, text, size=10)
            set_cell_border(cell, sz="8")
            if r == 0:
                shade_cell(cell, "F4F1EA")
        doc.add_paragraph()

    add_p(doc, "見證及收執確認", font="標楷體", size=13, bold=True, align="left", sb=8, sa=6, line=22)
    wt = doc.add_table(rows=2, cols=3)
    headers = ["受託仲介", "受託代書", "履約保證機構"]
    bodies = [
        f"{AGENCY}\n業務代表：{AGENT}\n\n簽署／店章：\n\n\n日期：　　年　　月　　日",
        "事務所：＿＿＿＿＿＿＿＿＿＿＿＿\n代書：＿＿＿＿＿＿＿＿＿＿\n\n簽署／職章：\n\n\n日期：　　年　　月　　日",
        "機構：＿＿＿＿＿＿＿＿＿＿＿＿\n承辦：＿＿＿＿＿＿＿＿＿＿\n\n簽署／職章：\n\n\n日期：　　年　　月　　日",
    ]
    for i, h in enumerate(headers):
        set_cell_text(wt.cell(0, i), h, font="標楷體", size=11, bold=True, align="center")
        shade_cell(wt.cell(0, i), "EFEFE9")
        set_cell_border(wt.cell(0, i))
        set_cell_text(wt.cell(1, i), bodies[i], size=10)
        set_cell_border(wt.cell(1, i))
        wt.cell(1, i).paragraphs[0].paragraph_format.space_after = Pt(48)

    doc.add_page_break()
    add_p(doc, "附件一　其他共有人價款分配明細表", font="標楷體", size=16, bold=True, align="center", sb=4, sa=8, line=26)
    add_p(
        doc,
        f"本表係第五條第二項之分配依據。可分配價金扣除立同意書人{ELDER}依第五條第一項取得之數額後，"
        "其餘額按下表分配。比例合計應為百分之一百。金額欄得於結算後由代書補填。",
        size=11,
        sa=8,
        line=20,
    )

    t1 = doc.add_table(rows=10, cols=4)
    heads = ["共有人姓名", "分配比例或金額", "預估／實際分配金額", "簽認"]
    for i, h in enumerate(heads):
        set_cell_text(t1.cell(0, i), h, font="標楷體", size=10, bold=True, align="center")
        shade_cell(t1.cell(0, i), "2A2A2A")
        for run in t1.cell(0, i).paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_border(t1.cell(0, i), color="222222")
    row_data = [[ELDER, "依第五條第一項（百分比或固定金額）", "依第五條第一項計算", ""]]
    for _ in range(7):
        row_data.append(["＿＿＿＿＿＿＿＿", "＿＿＿＿＿＿％　或　新臺幣＿＿＿＿＿＿＿＿元", "＿＿＿＿＿＿＿＿＿＿", ""])
    row_data.append(["合計", "100％（其餘額部分）", "可分配價金全額", ""])
    for r, row in enumerate(row_data, start=1):
        for c, val in enumerate(row):
            set_cell_text(t1.cell(r, c), val, size=10, align="center")
            set_cell_border(t1.cell(r, c))
            t1.cell(r, c).paragraphs[0].paragraph_format.space_before = Pt(8)
            t1.cell(r, c).paragraphs[0].paragraph_format.space_after = Pt(8)

    add_p(
        doc,
        "備註：如實際可分配價金與簽約時預估不同，除第五條第四項另有約定外，按本表比例調整。本表塗改處應由全體共有人蓋章。",
        size=10,
        sb=8,
        sa=6,
        line=16,
    )
    add_p(doc, "代書結算日期：中華民國＿＿＿＿年＿＿＿月＿＿＿日　　代書簽章：＿＿＿＿＿＿＿＿＿＿", size=12, sa=16)

    add_p(doc, "附件二　指定撥款帳戶及應備文件", font="標楷體", size=16, bold=True, align="center", sb=12, sa=8, line=26)
    add_p(
        doc,
        "請以共有人本人名義帳戶為原則。非本人帳戶應另附同意書。撥款前請核對戶名、銀行及帳號無誤。",
        size=11,
        sa=8,
        line=20,
    )
    t2 = doc.add_table(rows=9, cols=5)
    heads2 = ["戶名（共有人）", "金融機構／分行", "帳號", "身分證字號", "簽認"]
    for i, h in enumerate(heads2):
        set_cell_text(t2.cell(0, i), h, font="標楷體", size=10, bold=True, align="center")
        shade_cell(t2.cell(0, i), "2A2A2A")
        for run in t2.cell(0, i).paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_border(t2.cell(0, i))
    t2_rows = [[ELDER, "＿＿＿＿＿＿＿＿＿＿", "＿＿＿＿＿＿＿＿＿＿＿＿", "＿＿＿＿＿＿＿＿＿＿", ""]]
    for _ in range(7):
        t2_rows.append(["＿＿＿＿＿＿＿＿", "＿＿＿＿＿＿＿＿＿＿", "＿＿＿＿＿＿＿＿＿＿＿＿", "＿＿＿＿＿＿＿＿＿＿", ""])
    for r, row in enumerate(t2_rows, start=1):
        for c, val in enumerate(row):
            set_cell_text(t2.cell(r, c), val, size=10, align="center")
            set_cell_border(t2.cell(r, c))
            t2.cell(r, c).paragraphs[0].paragraph_format.space_before = Pt(8)
            t2.cell(r, c).paragraphs[0].paragraph_format.space_after = Pt(8)

    add_p(doc, "交割時請備文件（請代書勾選）", font="標楷體", size=13, bold=True, align="left", sb=14, sa=6, line=22)
    add_p(doc, "□　國民身分證正本及影本　　□　印鑑章及印鑑證明　　□　所有權狀正本", align="left")
    add_p(doc, "□　委託書／同意書正本　　□　戶口名簿或戶籍謄本　　□　指定帳戶存摺封面影本", align="left")
    add_p(doc, "□　他項權利塗銷相關文件（如有）　　□　其他：＿＿＿＿＿＿＿＿＿＿＿＿", align="left")

    add_p(
        doc,
        "本同意書係依立同意書人提供之協議意旨撰擬，供全體共有人簽署使用。"
        "涉及稅負、持分、優先購買權及繼承關係等事項，建議於用印前請受託代書或律師核閱。",
        size=9,
        align="center",
        sb=18,
        sa=4,
        line=14,
    )
    add_p(
        doc,
        f"標的門牌：{ADDRESS}　　受託仲介：{AGENCY}　業務代表：{AGENT}",
        size=9,
        align="center",
        sb=0,
        sa=4,
        line=14,
    )

    doc.save(DOCX_PATH)


def build_html() -> None:
    html = f"""<!DOCTYPE html>
<html lang="zh-Hant">
<head>
<meta charset="utf-8">
<title>房屋及土地出售價款分配同意書</title>
<style>
  @page {{ size: A4; margin: 16mm 16mm 18mm 16mm; }}
  * {{ box-sizing: border-box; }}
  body {{
    font-family: "I.Ming", "AR PL UMing TW", "Noto Serif CJK TC", "PMingLiU", "新細明體", serif;
    color: #1a1a1a;
    line-height: 1.75;
    max-width: 210mm;
    margin: 0 auto;
    padding: 18mm 16mm 20mm;
    background: #fff;
  }}
  h1 {{
    font-family: "AR PL UKai TW", "DFKai-SB", "標楷體", serif;
    font-size: 22pt;
    text-align: center;
    font-weight: normal;
    letter-spacing: 0.12em;
    margin: 0 0 6px;
  }}
  .rule {{ border: none; border-top: 2.2px solid #222; margin: 8px 0 1px; }}
  .rule-thin {{ border: none; border-top: 0.6px solid #222; margin: 0 0 14px; }}
  .sub, .meta {{ text-align: center; font-size: 11pt; margin: 2px 0; }}
  h2 {{
    font-family: "AR PL UKai TW", "DFKai-SB", "標楷體", serif;
    font-size: 13pt;
    font-weight: normal;
    margin: 1.15em 0 0.4em;
  }}
  p {{ font-size: 12pt; margin: 0.35em 0; text-align: justify; text-justify: inter-ideograph; }}
  .indent {{ padding-left: 1.2em; }}
  .subindent {{ padding-left: 2.2em; }}
  .label {{ font-family: "AR PL UKai TW", "標楷體", serif; font-size: 13pt; margin-top: 1em; }}
  .date {{ text-align: center; font-family: "AR PL UKai TW", "標楷體", serif; font-size: 13pt; margin-top: 1.4em; }}
  table {{ border-collapse: collapse; width: 100%; margin: 10px 0 16px; }}
  th, td {{ border: 1px solid #333; padding: 8px 6px; font-size: 10.5pt; vertical-align: middle; }}
  th {{ background: #2a2a2a; color: #fff; font-family: "AR PL UKai TW", "標楷體", serif; font-weight: normal; }}
  .sign {{ margin: 10px 0 14px; }}
  .sign td {{ height: 42px; }}
  .headrow td {{ background: #f4f1ea; }}
  .note {{ font-size: 10pt; color: #333; }}
  .center {{ text-align: center; }}
  .small {{ font-size: 9.5pt; color: #444; text-align: center; }}
  .page-break {{ page-break-before: always; }}
  @media print {{
    body {{ padding: 0; }}
    .page-break {{ page-break-before: always; }}
  }}
</style>
</head>
<body>
<h1>房屋及土地出售價款分配同意書</h1>
<hr class="rule"><hr class="rule-thin">
<p class="sub">（全體房屋所有權共有人簽署）</p>
<p class="meta">買賣標的門牌：{ADDRESS}</p>

<p class="label">立同意書人</p>
<p>本件買賣標的（定義如第一條）之<b>全體房屋及土地所有權共有人</b>（以下合稱「全體共有人」或「立同意書人」；各人之姓名、國民身分證統一編號、戶籍地址、權利範圍及指定撥款帳戶，詳如本同意書末之簽署欄及附件）。</p>
<p>緣立同意書人等為第一條所載房地之登記所有權人，為使出售程序明確、價款分配公允、交割撥款有據，經全體協商後，<b>本於自由意願</b>，訂立本同意書，以資共同遵守：</p>

<h2>第一條　買賣標的</h2>
<p>本同意書所稱「買賣標的」，係指下列房屋及其基地，並包含依法應隨同移轉之權利：</p>
<p class="indent">（一）門牌地址：<b>{ADDRESS}</b>。</p>
<p class="indent">（二）土地標示：＿＿＿＿縣（市）＿＿＿＿＿＿鄉（鎮、市、區）＿＿＿＿＿＿段＿＿＿＿小段＿＿＿＿＿＿地號；地目＿＿＿＿；面積＿＿＿＿＿＿＿＿平方公尺；權利範圍＿＿＿＿＿＿＿＿＿＿。</p>
<p class="indent">（三）建物標示：建號＿＿＿＿＿＿＿＿＿＿；門牌同第（一）款；建築完成日期＿＿＿＿＿＿＿＿＿＿；面積＿＿＿＿＿＿＿＿平方公尺（含附屬建物）；權利範圍＿＿＿＿＿＿＿＿＿＿。</p>
<p>前項土地及建物之詳細標示，如與地政機關最新登記謄本不符，以登記謄本為準；得於空白處補正或檢附謄本作為附件，視為本同意書之一部分。</p>

<h2>第二條　委託銷售</h2>
<p>全體共有人同意將買賣標的委託<b>{AGENCY}</b>辦理銷售，並指定業務代表<b>{AGENT}先生</b>執行帶看、議價、媒合及簽約協商等銷售業務。</p>
<p>委託期間、仲介服務報酬及其他權利義務，另以全體共有人與該加盟店簽訂之不動產委託銷售契約為準。該契約與本同意書抵觸時，<b>有關價款分配及撥款指示，以本同意書為準</b>。</p>

<h2>第三條　委託代書辦理交割</h2>
<p>全體共有人同意本件買賣之簽約、用印、稅費申報、所有權移轉登記、點交、抵押權塗銷（如有）及其他交割程序，委由代書依相關法令、買賣契約及本同意書辦理。</p>
<p class="indent">受託代書／事務所名稱：＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿　　聯絡電話：＿＿＿＿＿＿＿＿＿＿＿＿</p>
<p class="indent">履約保證機構名稱：＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿　　專案編號：＿＿＿＿＿＿＿＿＿＿＿＿</p>

<h2>第四條　出售價金及應扣除費用</h2>
<p class="indent">一、本次出售總價金，以全體共有人與買方簽訂之不動產買賣契約所載價款為準（下稱「成交總價」）。預定或已議定之成交總價：新臺幣＿＿＿＿＿＿＿＿＿＿＿＿＿＿元整（未填者，以日後簽訂之買賣契約為準）。</p>
<p class="indent">二、成交總價應優先扣除下列費用及稅捐後，以其餘額作為本同意書之可分配金額（下稱「可分配價金」）：</p>
<p class="subindent">（一）仲介服務報酬；</p>
<p class="subindent">（二）土地增值稅、依法應由出賣人負擔之房屋稅及地價稅、印花稅、房地合一所得稅，以及其他依法應由出賣人負擔之稅捐；</p>
<p class="subindent">（三）代書費、地政規費、謄本費、印鑑證明及印鑑登記相關費用、抵押權塗銷費、履約保證相關費用；</p>
<p class="subindent">（四）為完成本件交易必要，且經全體共有人同意或依法必須支出之其他費用。</p>
<p class="indent">三、前項費用之實際金額，以代書製作之結算明細、繳款收據或稅單為準。全體共有人簽署本同意書，即授權代書據實扣除後，依第五條及附件分配撥款。</p>

<h2>第五條　價款分配方式</h2>
<p class="indent">一、立同意書人<b>{ELDER}</b>（於兄弟姊妹間為最年長之兄長），基於親屬間協商之結果，同意自可分配價金中領取下列方式之一（請擇一勾選並以正楷填載；勾選後之空白處如未填載，本項不生分配效力，應由全體共有人補簽後始得撥款）：</p>
<p class="indent">□　可分配價金之百分之＿＿＿＿＿＿（即　＿＿＿＿　％；大寫：百分之＿＿＿＿＿＿＿＿）。</p>
<p class="indent">□　固定金額新臺幣＿＿＿＿＿＿＿＿＿＿＿＿＿＿元整（大寫：新臺幣＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿元整）。</p>
<p class="indent">二、可分配價金扣除{ELDER}依前項取得之數額後，<b>其餘價款由其他所有權共有人協議分配</b>，其姓名、比例或金額，詳如附件一「其他共有人價款分配明細表」。</p>
<p class="indent">三、全體共有人確認並同意：本條約定之分配方式，係本於兄弟姊妹及共有人間之特別合意，<b>得與地政登記之應有部分（持分）比例不相同</b>；一經簽署，不得再以登記持分、出資多寡、管理貢獻或其他理由異議，或請求重新分配。</p>
<p class="indent">四、如成交總價或應扣除費用於簽約後有增減，致可分配價金變動時：勾選「百分比」者，應按原約定比例自動調整；勾選「固定金額」者，{ELDER}所領固定金額不變，其餘共有人就剩餘款項再協議，並以書面補正附件一。</p>

<h2>第六條　撥款指示</h2>
<p class="indent">一、全體共有人在此同意並指示受託代書及履約保證機構，於買賣價金依約可撥付時，<b>應逕依本同意書第五條、附件一及附件二辦理分配撥款</b>，無需再逐一徵詢各共有人。各方不得對依本同意書所為之撥款提出異議。</p>
<p class="indent">二、各共有人指定之撥款帳戶詳如簽署欄及附件二。帳戶資料如有變更，應於撥款日至少三個營業日前，以書面通知代書及履約保證機構，並由變更人簽名或蓋章。</p>
<p class="indent">三、因帳戶資料錯誤、未及時變更、拒絕提供證件印鑑，或受款帳戶非其本人名義，致撥款遲延、退匯或誤入他人帳戶者，由該共有人自行負責，不得因此拒絕交割或主張本同意書無效。</p>
<p class="indent">四、代書及履約保證機構依本同意書完成撥款後，就該部分價金對全體共有人之給付義務即為清償完畢。</p>

<h2>第七條　聲明、保證及效力</h2>
<p class="indent">一、立同意書人簽署本同意書，均出於自由意願，並已詳細閱讀、充分瞭解全部內容，並願受其拘束。</p>
<p class="indent">二、立同意書人保證其為買賣標的之合法共有人（或有權處分之人），並已將所知之其他共有人、他項權利、優先購買權、租賃、占用或其他可能影響交易之事由，據實告知受託代書及仲介。</p>
<p class="indent">三、如因立同意書人隱匿、虛偽陳述或未配合提供權狀、印鑑證明、身分證件，致本件交易或他方受有損害，應由可歸責之立同意書人負損害賠償責任。</p>
<p class="indent">四、本同意書如有塗改，應於塗改處由全體共有人蓋章，始生效力。空白未填部分，不得擅自添註。</p>

<h2>第八條　補充約定、份數、準據法及生效</h2>
<p class="indent">一、本同意書未盡事宜，依中華民國民法、土地法、土地登記規則及其他相關法令辦理；必要時得由全體共有人另行簽立書面協議補充，該書面與本同意書有同等效力。</p>
<p class="indent">二、本同意書之解釋、效力與履行，以中華民國法律為準據法。如因此發生訴訟，同意以臺灣宜蘭地方法院為第一審管轄法院。</p>
<p class="indent">三、本同意書壹式　＿＿＿＿　份，由全體共有人、受託代書、履約保證機構及受託仲介店各執乙份為憑；影本或掃描檔經核與正本無異者，與正本有同一效力。</p>
<p class="indent">四、本同意書於全體共有人親自簽名或蓋用與印鑑證明相符之印鑑章之日起生效。</p>

<p class="date">此　　據</p>
<p class="date">中華民國　　　　年　　　　月　　　　日</p>

<div class="page-break"></div>
<h1>立同意書人簽署欄</h1>
<p class="small">以下簽署人確為本同意書之立同意書人，已閱讀並同意全文及附件，願依約履行。未使用之簽署欄，請劃記「空白」並由在場共有人蓋章註銷。</p>
"""

    def sign_table(name, note):
        display = f"{name}（{note}）" if name else f"{'＿'*12}（{note}）"
        return f"""
<table class="sign">
  <tr class="headrow">
    <td>姓名：{display}</td>
    <td>國民身分證統一編號：＿＿＿＿＿＿＿＿＿＿＿＿＿＿</td>
  </tr>
  <tr>
    <td>戶籍地址：＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿</td>
    <td>聯絡電話：＿＿＿＿＿＿＿＿＿＿＿＿＿＿</td>
  </tr>
  <tr>
    <td>權利範圍（持分）：＿＿＿＿＿＿＿＿＿＿＿＿</td>
    <td>簽名或蓋章：</td>
  </tr>
  <tr>
    <td>指定撥款帳戶（銀行／帳號／戶名）：＿＿＿＿＿＿＿＿＿＿＿＿</td>
    <td>（請蓋印鑑章，與印鑑證明相符）</td>
  </tr>
</table>
"""

    html += sign_table(ELDER, "長兄／共有人")
    for i in range(2, 9):
        html += sign_table("", f"共有人 {i}")

    html += f"""
<h2>見證及收執確認</h2>
<table>
  <tr>
    <th>受託仲介</th>
    <th>受託代書</th>
    <th>履約保證機構</th>
  </tr>
  <tr>
    <td>{AGENCY}<br>業務代表：{AGENT}<br><br>簽署／店章：<br><br><br>日期：　　年　　月　　日</td>
    <td>事務所：＿＿＿＿＿＿＿＿＿＿＿＿<br>代書：＿＿＿＿＿＿＿＿＿＿<br><br>簽署／職章：<br><br><br>日期：　　年　　月　　日</td>
    <td>機構：＿＿＿＿＿＿＿＿＿＿＿＿<br>承辦：＿＿＿＿＿＿＿＿＿＿<br><br>簽署／職章：<br><br><br>日期：　　年　　月　　日</td>
  </tr>
</table>

<div class="page-break"></div>
<h1>附件一　其他共有人價款分配明細表</h1>
<p>本表係第五條第二項之分配依據。可分配價金扣除立同意書人{ELDER}依第五條第一項取得之數額後，其餘額按下表分配。比例合計應為百分之一百。金額欄得於結算後由代書補填。</p>
<table>
  <tr>
    <th>共有人姓名</th>
    <th>分配比例或金額</th>
    <th>預估／實際分配金額</th>
    <th>簽認</th>
  </tr>
  <tr>
    <td class="center">{ELDER}</td>
    <td class="center">依第五條第一項（百分比或固定金額）</td>
    <td class="center">依第五條第一項計算</td>
    <td></td>
  </tr>
"""
    for _ in range(7):
        html += """  <tr>
    <td class="center">＿＿＿＿＿＿＿＿</td>
    <td class="center">＿＿＿＿＿＿％　或　新臺幣＿＿＿＿＿＿＿＿元</td>
    <td class="center">＿＿＿＿＿＿＿＿＿＿</td>
    <td></td>
  </tr>
"""
    html += f"""  <tr>
    <td class="center">合計</td>
    <td class="center">100％（其餘額部分）</td>
    <td class="center">可分配價金全額</td>
    <td></td>
  </tr>
</table>
<p class="note">備註：如實際可分配價金與簽約時預估不同，除第五條第四項另有約定外，按本表比例調整。本表塗改處應由全體共有人蓋章。</p>
<p>代書結算日期：中華民國＿＿＿＿年＿＿＿月＿＿＿日　　代書簽章：＿＿＿＿＿＿＿＿＿＿</p>

<h1>附件二　指定撥款帳戶及應備文件</h1>
<p>請以共有人本人名義帳戶為原則。非本人帳戶應另附同意書。撥款前請核對戶名、銀行及帳號無誤。</p>
<table>
  <tr>
    <th>戶名（共有人）</th>
    <th>金融機構／分行</th>
    <th>帳號</th>
    <th>身分證字號</th>
    <th>簽認</th>
  </tr>
  <tr>
    <td class="center">{ELDER}</td>
    <td></td><td></td><td></td><td></td>
  </tr>
"""
    for _ in range(7):
        html += "  <tr><td></td><td></td><td></td><td></td><td></td></tr>\n"
    html += f"""</table>
<h2>交割時請備文件（請代書勾選）</h2>
<p>□　國民身分證正本及影本　　□　印鑑章及印鑑證明　　□　所有權狀正本</p>
<p>□　委託書／同意書正本　　□　戶口名簿或戶籍謄本　　□　指定帳戶存摺封面影本</p>
<p>□　他項權利塗銷相關文件（如有）　　□　其他：＿＿＿＿＿＿＿＿＿＿＿＿</p>
<hr class="rule-thin">
<p class="small">本同意書係依立同意書人提供之協議意旨撰擬，供全體共有人簽署使用。涉及稅負、持分、優先購買權及繼承關係等事項，建議於用印前請受託代書或律師核閱。</p>
<p class="small">標的門牌：{ADDRESS}　　受託仲介：{AGENCY}　業務代表：{AGENT}</p>
</body>
</html>
"""
    with open(HTML_PATH, "w", encoding="utf-8") as f:
        f.write(html)


def main() -> None:
    build_pdf()
    build_docx()
    build_html()
    print("Wrote:")
    print(" ", PDF_PATH)
    print(" ", DOCX_PATH)
    print(" ", HTML_PATH)


if __name__ == "__main__":
    main()
